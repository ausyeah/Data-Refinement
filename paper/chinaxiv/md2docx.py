# -*- coding: utf-8 -*-
"""预印本 Markdown -> ChinaXiv 规范 DOCX 排版转换器（参数化版）。

版式依据《中国科技论文预发布平台论文提交格式规范-2016-02-01》中文论文样式表：
A4；文题 2号宋体(22pt)加粗居中；作者/单位 5号宋体居中；
摘要 5号宋体五段式；关键词 5号；正文 小4宋体(12pt)；1级标题 4号宋体(14pt)粗体；
2级标题 小4楷体_GB2312；表格 Table Grid 五号；参考文献五号；页边距 2.5cm。

用法：python md2docx.py --in 源.md --out 输出.docx
"""
import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

SONG, HEI, KAI, TNR = "宋体", "黑体", "楷体_GB2312", "Times New Roman"


def set_run(run, *, ea=SONG, ascii_=TNR, size=12, bold=False):
    run.font.name = ascii_
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), ea)


def add_runs(par, text, *, ea=SONG, ascii_=TNR, size=12, base_bold=False):
    for i, seg in enumerate(re.split(r"\*\*", text)):
        if seg == "":
            continue
        run = par.add_run(seg)
        set_run(run, ea=ea, ascii_=ascii_, size=size, bold=base_bold or i % 2 == 1)


def para(doc, text="", *, ea=SONG, ascii_=TNR, size=12, bold=False, align=None,
         indent_chars=0, line=1.5, before=0, after=0, left_cm=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = line
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if align is not None:
        pf.alignment = align
    if indent_chars:
        pf.first_line_indent = Pt(size * indent_chars)
    if left_cm is not None:
        pf.left_indent = Cm(left_cm)
    if text:
        add_runs(p, text, ea=ea, ascii_=ascii_, size=size, base_bold=bold)
    return p


def merge_soft_wraps(lines):
    CJK = lambda ch: '\u4e00' <= ch <= '\u9fff'
    out, buf = [], []

    def flush():
        if buf:
            joined = buf[0]
            for nxt in buf[1:]:
                joined += ("" if CJK(joined[-1]) or CJK(nxt[0]) else " ") + nxt
            out.append(joined.strip())
            buf.clear()
    for ln in lines:
        s = ln.strip()
        if (s == "" or s == "---" or s.startswith(("#", "|", ">", "- "))
                or re.match(r"\[\d+\]", s)):
            flush()
            out.append(ln)
        else:
            buf.append(s)
    flush()
    return out


def strip_inline_md(lines):
    """去反引号、单星斜体标记与 Markdown 转义反斜杠（保留 ** 粗体由 add_runs 处理）。"""
    out = []
    for ln in lines:
        ln = re.sub(r"`([^`]*)`", r"\1", ln)
        ln = re.sub(r"\\([\*'#])", r"\1", ln)                 # 去除 \* \' 等转义
        ln = re.sub(r"(?<!\*)\*([^*\n]+)\*(?!\*)", r"\1", ln) # 单星斜体 → 去星
        out.append(ln)
    return out


def build(src: Path, out: Path):
    raw = src.read_text(encoding="utf-8").splitlines()
    raw = strip_inline_md(raw)
    lines = merge_soft_wraps(raw)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2.5)
    # 页脚居中页码
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = fp.add_run()
    fld = r1._element
    import docx.oxml.ns as ns
    b, i_, e = ns.qn("w:fldChar"), ns.qn("w:instrText"), ns.qn("w:fldChar")
    fldB = fld.makeelement(b, {ns.qn("w:fldCharType"): "begin"}); fld.append(fldB)
    it = fld.makeelement(i_, {}); it.text = " PAGE "; fld.append(it)
    fldE = fld.makeelement(e, {ns.qn("w:fldCharType"): "end"}); fld.append(fldE)
    for r in fp.runs: r.font.size = Pt(9)

    h1_seen = 0
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        s = ln.strip()

        if s == "" or s == "---":
            i += 1
            continue

        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            ncols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=ncols)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r, row in enumerate(rows):
                for c in range(ncols):
                    cell = table.cell(r, c)
                    cell.paragraphs[0].paragraph_format.line_spacing = 1.0
                    text = row[c] if c < len(row) else ""
                    add_runs(cell.paragraphs[0], text, size=10.5, base_bold=(r == 0))
            para(doc, "", size=6)
            continue

        if s.startswith("###"):
            para(doc, s.lstrip("#").strip(), ea=SONG, size=12, bold=True,
                 indent_chars=2, before=8, after=4)
        elif s.startswith("##"):
            t = s.lstrip("#").strip()
            if re.search(r"[\u4e00-\u9fff]", t):
                para(doc, t, ea=KAI, size=12, bold=False, before=10, after=4)
            else:                                   # 英文二级标题：TNR 粗体，不再用楷体
                para(doc, t, ea=HEI, ascii_=TNR, size=12, bold=True, before=10, after=4)
        elif s.startswith("#"):
            text = s.lstrip("#").strip()
            h1_seen += 1
            # 不自动分页（连续排版）；摘要/Abstract 用规范样式
            if text == "摘要" or text == "Abstract":
                para(doc, text, ea=HEI, ascii_=TNR, size=10.5, bold=True,
                     align=WD_ALIGN_PARAGRAPH.CENTER, before=10, after=6)
            elif h1_seen == 1:
                para(doc, text, ea=SONG, size=22, bold=True,
                     align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=10)
            else:
                para(doc, text, ea=SONG, size=14, bold=True,
                     align=WD_ALIGN_PARAGRAPH.LEFT, before=6, after=8)
        elif s.startswith(">"):
            para(doc, s.lstrip("> ").strip(), ea=KAI, size=10.5, left_cm=0.74,
                 line=1.25)
        elif s.startswith("- "):
            p = para(doc, s[2:].strip(), size=12)
            p.paragraph_format.left_indent = Cm(0.74)
        else:
            text = s
            align = None
            indent = 2
            if re.fullmatch(r"\*\*(表|图)[\d\s．.\-＋+－−～~、（）()：:*×].*\*\*", s):
                align = WD_ALIGN_PARAGRAPH.CENTER
                indent = 0
            if s.startswith("作者：") or s.startswith("单位："):
                align = WD_ALIGN_PARAGRAPH.CENTER
                indent = 0
                size = 10.5
                para(doc, text, size=size, align=align)
                i += 1
                continue
            if re.match(r"\[\d+\]", s):
                para(doc, text, size=10.5, line=1.25)
                i += 1
                continue
            ABSTRACT_PFX = ("**目的", "**方法", "**结果", "**局限", "**结论",
                            "**关键词", "**Objective", "**Methods", "**Results",
                            "**Limitations", "**Conclusions", "**Keywords")
            if s.startswith(ABSTRACT_PFX):      # 摘要/关键词段：五号 10.5
                para(doc, text, size=10.5)
                i += 1
                continue
            para(doc, text, size=12, align=align, indent_chars=indent)
        i += 1

    doc.save(out)
    print(f"[OK] {out}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--in", dest="src", required=True)
    ap.add_argument("--out", dest="out", required=True)
    args = ap.parse_args()
    build(Path(args.src), Path(args.out))
